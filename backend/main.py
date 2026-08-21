"""FastAPI backend for LLM Council."""

import os
from fastapi import FastAPI, HTTPException, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
import uuid
import json
import asyncio

from . import storage
from .council import (
    run_full_council,
    generate_conversation_title,
    stage1_collect_responses,
    stage2_collect_rankings,
    stage3_synthesize_final,
    calculate_aggregate_rankings,
    run_hybrid_council,
    hybrid_phase1_socratic,
    hybrid_phase2_debate,
    hybrid_phase3_devils_advocate,
    hybrid_phase4_synthesis,
    run_consensus_debate_stream,
)
from .freemodels import get_free_models

app = FastAPI(title="LLM Council API")

# Enable CORS for local development
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class CreateConversationRequest(BaseModel):
    """Request to create a new conversation."""
    pass


class SendMessageRequest(BaseModel):
    """Request to send a message in a conversation."""
    content: str
    mode: str = "council"  # "council" | "hybrid" | "consensus"
    # Consensus mode options:
    model_ids: List[str] = []
    chairman_id: Optional[str] = None
    max_rounds: int = 4


class RenameConversationRequest(BaseModel):
    """Request to rename a conversation."""
    title: str


class ConversationMetadata(BaseModel):
    """Conversation metadata for list view."""
    id: str
    created_at: str
    title: str
    message_count: int


class Conversation(BaseModel):
    """Full conversation with all messages."""
    id: str
    created_at: str
    title: str
    messages: List[Dict[str, Any]]


@app.get("/")
async def root():
    """Health check endpoint."""
    return {"status": "ok", "service": "LLM Council API"}


@app.get("/api/ping")
async def ping_models():
    """
    Ping all configured LLM models with a minimal prompt.
    Returns Server-Sent Events as each model responds.
    """
    import time as _time
    from .config import COUNCIL_MODELS, CHAIRMAN_CONFIG, DEVILS_ADVOCATE_CONFIG
    from .providers import query_model

    # Build deduplicated list of all models to test
    seen = set()
    models_to_test = []
    for config in COUNCIL_MODELS:
        key = (id(config["provider"]), config["model"])
        if key not in seen:
            seen.add(key)
            models_to_test.append(config)
    for extra in [CHAIRMAN_CONFIG, DEVILS_ADVOCATE_CONFIG]:
        key = (id(extra["provider"]), extra["model"])
        if key not in seen:
            seen.add(key)
            models_to_test.append(extra)

    total = len(models_to_test)
    ping_message = [{"role": "user", "content": "Reply with only the word: pong"}]

    async def ping_one(config):
        name = config["name"]
        start = _time.perf_counter()
        try:
            response = await query_model(
                config["provider"],
                config["model"],
                ping_message,
                timeout=120.0
            )
            elapsed = round((_time.perf_counter() - start) * 1000)
            if response and response.get("content"):
                return {"model": name, "status": "ok", "latency_ms": elapsed, "response": response["content"].strip()[:50]}
            else:
                return {"model": name, "status": "error", "latency_ms": elapsed, "error": "No response"}
        except Exception as e:
            elapsed = round((_time.perf_counter() - start) * 1000)
            return {"model": name, "status": "error", "latency_ms": elapsed, "error": str(e)[:100]}

    async def event_generator():
        # Send initial event with model count
        yield f"data: {json.dumps({'type': 'ping_start', 'total': total, 'models': [c['name'] for c in models_to_test]})}\n\n"

        # Fire all pings concurrently, yield as each completes
        tasks = {asyncio.create_task(ping_one(c)): c["name"] for c in models_to_test}
        for coro in asyncio.as_completed(tasks):
            result = await coro
            yield f"data: {json.dumps({'type': 'ping_result', 'data': result})}\n\n"

        yield f"data: {json.dumps({'type': 'ping_complete'})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


def _extract_text_from_file(ext: str, data: bytes, filename: str) -> str:
    """Extract plain text from an uploaded file based on its extension."""
    import io

    if ext in {"txt", "sh", "py", "md"}:
        return data.decode("utf-8", errors="replace")

    elif ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages) if pages else "(No text could be extracted from this PDF)"
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse PDF: {e}")

    elif ext == "docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs) if paragraphs else "(No text could be extracted from this DOCX)"
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse DOCX: {e}")

    elif ext == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_str = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_str.strip():
                        rows.append(row_str)
            return "\n".join(rows) if rows else "(No data found in XLSX)"
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse XLSX: {e}")

    elif ext == "xls":
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=data)
            rows = []
            for sheet in wb.sheets():
                rows.append(f"[Sheet: {sheet.name}]")
                for rx in range(sheet.nrows):
                    row_str = "\t".join(str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols))
                    if row_str.strip():
                        rows.append(row_str)
            return "\n".join(rows) if rows else "(No data found in XLS)"
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse XLS: {e}")

    raise HTTPException(status_code=400, detail=f"Unsupported file type: .{ext}")


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    """
    Upload a file and extract its text content.
    Supports: pdf, docx, txt, sh, py, md, xls, xlsx
    Returns: { "text": str, "filename": str, "size": int }
    """
    ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "sh", "py", "md", "xls", "xlsx"}
    MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB

    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '.{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    data = await file.read()
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File too large. Maximum size is 20 MB.")

    text = _extract_text_from_file(ext, data, filename)

    return {
        "text": text,
        "filename": filename,
        "size": len(data)
    }


@app.get("/api/models/free")
async def list_free_models():
    """
    List all currently-free OpenRouter chat models from the models.dev catalog.
    These models cost nothing but require OPENROUTER_API_KEY (free tier).
    """
    if not os.getenv("OPENROUTER_API_KEY"):
        return {
            "requires_key": True,
            "models": [],
            "error": "OPENROUTER_API_KEY is not configured. Add it to your .env file (the OpenRouter free tier costs nothing).",
        }
    try:
        models = await get_free_models()
        return {"requires_key": False, "models": models, "error": None}
    except Exception as e:
        return {"requires_key": False, "models": [], "error": f"Could not fetch model catalog: {e}"}


@app.get("/api/conversations", response_model=List[ConversationMetadata])
async def list_conversations():
    """List all conversations (metadata only)."""
    return storage.list_conversations()


@app.post("/api/conversations", response_model=Conversation)
async def create_conversation(request: CreateConversationRequest):
    """Create a new conversation."""
    conversation_id = str(uuid.uuid4())
    conversation = storage.create_conversation(conversation_id)
    return conversation


@app.get("/api/conversations/{conversation_id}", response_model=Conversation)
async def get_conversation(conversation_id: str):
    """Get a specific conversation with all its messages."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")
    return conversation


@app.post("/api/conversations/{conversation_id}/message")
async def send_message(conversation_id: str, request: SendMessageRequest):
    """
    Send a message and run the 3-stage council process.
    Returns the complete response with all stages.
    """
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    is_first_message = len(conversation["messages"]) == 0

    storage.add_user_message(conversation_id, request.content)

    if is_first_message:
        title = await generate_conversation_title(request.content)
        storage.update_conversation_title(conversation_id, title)

    stage1_results, stage2_results, stage3_result, metadata = await run_full_council(
        request.content
    )

    storage.add_assistant_message(
        conversation_id,
        stage1_results,
        stage2_results,
        stage3_result
    )

    return {
        "stage1": stage1_results,
        "stage2": stage2_results,
        "stage3": stage3_result,
        "metadata": metadata
    }


@app.post("/api/conversations/{conversation_id}/message/stream")
async def send_message_stream(conversation_id: str, request: SendMessageRequest):
    """
    Send a message and stream the 3-stage council process.
    Returns Server-Sent Events as each stage completes.
    """
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    is_first_message = len(conversation["messages"]) == 0

    async def event_generator():
        try:
            storage.add_user_message(conversation_id, request.content)

            title_task = None
            if is_first_message:
                title_task = asyncio.create_task(generate_conversation_title(request.content))

            yield f"data: {json.dumps({'type': 'stage1_start'})}\n\n"
            stage1_results = await stage1_collect_responses(request.content)
            yield f"data: {json.dumps({'type': 'stage1_complete', 'data': stage1_results})}\n\n"

            yield f"data: {json.dumps({'type': 'stage2_start'})}\n\n"
            stage2_results, label_to_model = await stage2_collect_rankings(request.content, stage1_results)
            aggregate_rankings = calculate_aggregate_rankings(stage2_results, label_to_model)
            yield f"data: {json.dumps({'type': 'stage2_complete', 'data': stage2_results, 'metadata': {'label_to_model': label_to_model, 'aggregate_rankings': aggregate_rankings}})}\n\n"

            yield f"data: {json.dumps({'type': 'stage3_start'})}\n\n"
            stage3_result = await stage3_synthesize_final(request.content, stage1_results, stage2_results)
            yield f"data: {json.dumps({'type': 'stage3_complete', 'data': stage3_result})}\n\n"

            if title_task:
                title = await title_task
                storage.update_conversation_title(conversation_id, title)
                yield f"data: {json.dumps({'type': 'title_complete', 'data': {'title': title}})}\n\n"

            storage.add_assistant_message(
                conversation_id,
                stage1_results,
                stage2_results,
                stage3_result
            )

            yield f"data: {json.dumps({'type': 'complete'})}\n\n"

        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


@app.post("/api/conversations/{conversation_id}/message/stream/hybrid")
async def send_message_stream_hybrid(conversation_id: str, request: SendMessageRequest):
    """
    Send a message and stream the 4-phase debate council process.
    Phase 1: Socratic | Phase 2: Debate | Phase 3: Devil's Advocate | Phase 4: Synthesis
    """
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    is_first_message = len(conversation["messages"]) == 0

    async def event_generator():
        try:
            storage.add_user_message(conversation_id, request.content)

            title_task = None
            if is_first_message:
                title_task = asyncio.create_task(generate_conversation_title(request.content))

            yield f"data: {json.dumps({'type': 'hybrid_phase1_start'})}\n\n"
            phase1_results = await hybrid_phase1_socratic(request.content)
            yield f"data: {json.dumps({'type': 'hybrid_phase1_complete', 'data': phase1_results})}\n\n"

            yield f"data: {json.dumps({'type': 'hybrid_phase2_start'})}\n\n"
            phase2_results = await hybrid_phase2_debate(request.content, phase1_results)
            yield f"data: {json.dumps({'type': 'hybrid_phase2_complete', 'data': phase2_results})}\n\n"

            yield f"data: {json.dumps({'type': 'hybrid_phase3_start'})}\n\n"
            phase3_result = await hybrid_phase3_devils_advocate(request.content, phase1_results, phase2_results)
            yield f"data: {json.dumps({'type': 'hybrid_phase3_complete', 'data': phase3_result})}\n\n"

            yield f"data: {json.dumps({'type': 'hybrid_phase4_start'})}\n\n"
            phase4_result = await hybrid_phase4_synthesis(request.content, phase1_results, phase2_results, phase3_result)
            yield f"data: {json.dumps({'type': 'hybrid_phase4_complete', 'data': phase4_result})}\n\n"

            if title_task:
                title = await title_task
                storage.update_conversation_title(conversation_id, title)
                yield f"data: {json.dumps({'type': 'title_complete', 'data': {'title': title}})}\n\n"

            hybrid_message = {
                "role": "assistant",
                "mode": "hybrid",
                "hybrid_phase1": phase1_results,
                "hybrid_phase2": phase2_results,
                "hybrid_phase3": phase3_result,
                "hybrid_phase4": phase4_result,
                "stage1": [],
                "stage2": [],
                "stage3": None,
                "metadata": {"mode": "hybrid"}
            }
            conversation = storage.get_conversation(conversation_id)
            conversation["messages"].append(hybrid_message)
            storage.save_conversation(conversation)

            yield f"data: {json.dumps({'type': 'complete'})}\n\n"

        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


@app.post("/api/conversations/{conversation_id}/message/stream/consensus")
async def send_message_stream_consensus(conversation_id: str, request: SendMessageRequest):
    """
    Send a message and stream a consensus debate between free OpenRouter models.
    The debate runs in rounds with majority voting; a chairman model (one of the
    participants, abstaining from votes) verifies alignment and synthesizes.
    """
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    if len(request.model_ids) < 3:
        raise HTTPException(status_code=400, detail="Consensus debates need at least 3 participating models.")
    if not request.chairman_id:
        raise HTTPException(status_code=400, detail="A chairman model must be designated.")
    if request.chairman_id not in request.model_ids:
        raise HTTPException(status_code=400, detail="The chairman must be one of the selected participants.")

    is_first_message = len(conversation["messages"]) == 0

    async def event_generator():
        consensus_message = {
            "role": "assistant",
            "mode": "consensus",
            "participants": [],
            "chairman": None,
            "rounds": [],
            "synthesis": None,
            "stage1": [],
            "stage2": [],
            "stage3": None,
            "hybrid_phase1": [],
            "hybrid_phase2": [],
            "hybrid_phase3": None,
            "hybrid_phase4": None,
            "metadata": {"mode": "consensus"},
        }

        try:
            storage.add_user_message(conversation_id, request.content)

            title_task = None
            if is_first_message:
                title_task = asyncio.create_task(generate_conversation_title(request.content))

            async for event in run_consensus_debate_stream(
                request.content,
                request.model_ids,
                request.chairman_id,
                max_rounds=request.max_rounds,
            ):
                etype = event["type"]
                data = event.get("data", {})

                if etype == "consensus_start":
                    consensus_message["participants"] = data["participants"]
                    consensus_message["chairman"] = data["chairman"]
                elif etype == "consensus_round_start":
                    consensus_message["rounds"].append({
                        "round": data["round"],
                        "statements": [],
                        "votes": {"yes": [], "no": []},
                        "reached": False,
                    })
                elif etype == "consensus_model_complete":
                    if consensus_message["rounds"]:
                        consensus_message["rounds"][-1]["statements"].append(data["statement"])
                elif etype == "consensus_round_complete":
                    if consensus_message["rounds"]:
                        entry = consensus_message["rounds"][-1]
                        entry["votes"] = data.get("votes", entry["votes"])
                        entry["reached"] = data.get("reached", False)
                elif etype == "consensus_chairman_review":
                    if consensus_message["rounds"]:
                        consensus_message["rounds"][-1]["review"] = data
                elif etype == "consensus_synthesis_complete":
                    consensus_message["synthesis"] = data

                yield f"data: {json.dumps({'type': etype, 'data': data})}\n\n"

            # Persist the finished message
            current = storage.get_conversation(conversation_id)
            current["messages"].append(consensus_message)
            storage.save_conversation(current)

            if title_task:
                title = await title_task
                storage.update_conversation_title(conversation_id, title)
                yield f"data: {json.dumps({'type': 'title_complete', 'data': {'title': title}})}\n\n"

            yield f"data: {json.dumps({'type': 'complete'})}\n\n"

        except Exception as e:
            print(f"[consensus] Stream error: {e}")
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


@app.delete("/api/conversations/{conversation_id}/messages")
async def clear_messages(conversation_id: str):
    """Clear all messages from a conversation but keep the conversation."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    conversation["messages"] = []
    storage.save_conversation(conversation)

    return {"success": True, "message": "Messages cleared"}


@app.put("/api/conversations/{conversation_id}/title")
async def rename_conversation(conversation_id: str, request: RenameConversationRequest):
    """Rename a conversation."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    storage.update_conversation_title(conversation_id, request.title)

    return {"success": True, "title": request.title}


@app.delete("/api/conversations/{conversation_id}")
async def delete_conversation(conversation_id: str):
    """Delete a conversation entirely."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    storage.delete_conversation(conversation_id)

    return {"success": True, "message": "Conversation deleted"}


@app.get("/api/conversations/{conversation_id}/export")
async def export_conversation(conversation_id: str):
    """Export a conversation as markdown."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    markdown = f"# {conversation['title']}\n\n"
    markdown += f"*Created: {conversation['created_at']}*\n\n"
    markdown += "---\n\n"

    for message in conversation["messages"]:
        if message["role"] == "user":
            # Parse optional file attachment from prepended block
            content = message["content"]
            import re as _re
            file_match = _re.match(r"^\[File: (.+?)\]", content)
            file_name = file_match.group(1) if file_match else None
            question_match = _re.search(r"\nUser question: ([\s\S]*)$", content)
            display_text = question_match.group(1) if question_match else content

            markdown += "## User\n\n"
            if file_name:
                markdown += f"📄 **Attached file:** `{file_name}`\n\n"
            markdown += f"{display_text}\n\n"

        elif message["role"] == "assistant":
            if message.get("mode") == "consensus":
                markdown += "## Consensus Mode Debate\n\n"

                participants = message.get("participants", [])
                chairman = message.get("chairman") or {}
                if participants:
                    names = ", ".join(p["name"] for p in participants)
                    markdown += f"**Participants:** {names}\n\n"
                if chairman:
                    markdown += f"**Chairman (non-voting):** {chairman.get('name', '')}\n\n"

                for round_entry in message.get("rounds", []):
                    round_label = f"Round {round_entry['round']}"
                    if round_entry.get("reached"):
                        round_label += " — Majority Consensus"
                    markdown += f"### {round_label}\n\n"

                    votes = round_entry.get("votes") or {}
                    if votes.get("yes") or votes.get("no"):
                        markdown += f"*Votes — YES: {', '.join(votes.get('yes', [])) or 'none'} · NO: {', '.join(votes.get('no', [])) or 'none'}*\n\n"

                    review = round_entry.get("review")
                    if review:
                        verdict = "ALIGNED" if review.get("aligned") else "CONFLICTED"
                        markdown += f"**Chairman review ({review.get('reviewer', 'Chairman')}):** {verdict} — {review.get('reasoning', '')}\n\n"

                    for stmt in round_entry.get("statements", []):
                        markdown += f"**{stmt['model']}:**\n\n{stmt['response']}\n\n"
                        if stmt.get("position"):
                            markdown += f"> **Final position:** {stmt['position']}\n\n"

                synthesis = message.get("synthesis") or {}
                if synthesis.get("response"):
                    markdown += "### Council Consensus (Final Synthesis)\n\n"
                    markdown += f"**{synthesis['model']}:**\n\n{synthesis['response']}\n\n"

            elif message.get("mode") == "hybrid":
                markdown += "## Debate Mode Council Response\n\n"

                markdown += "### Phase 1: Socratic (Initial Responses)\n\n"
                for response in message.get("hybrid_phase1", []):
                    markdown += f"**{response['model']}:**\n\n{response['response']}\n\n"

                markdown += "### Phase 2: Debate\n\n"
                for response in message.get("hybrid_phase2", []):
                    markdown += f"**{response['model']}:**\n\n{response['response']}\n\n"

                markdown += "### Phase 3: Devil's Advocate\n\n"
                p3 = message.get("hybrid_phase3") or {}
                if p3.get("response"):
                    markdown += f"**{p3['model']}:**\n\n{p3['response']}\n\n"

                markdown += "### Phase 4: Final Synthesis\n\n"
                p4 = message.get("hybrid_phase4") or {}
                if p4.get("response"):
                    markdown += f"**{p4['model']}:**\n\n{p4['response']}\n\n"

            else:
                markdown += "## LLM Council Response\n\n"

                markdown += "### Stage 1: Individual Responses\n\n"
                for response in message.get("stage1", []):
                    markdown += f"**{response['model']}:**\n\n{response['response']}\n\n"

                markdown += "### Stage 2: Peer Rankings\n\n"
                for ranking in message.get("stage2", []):
                    markdown += f"**{ranking['model']}:**\n\n{ranking['ranking']}\n\n"

                markdown += "### Stage 3: Final Synthesis\n\n"
                stage3 = message.get("stage3") or {}
                if stage3.get("response"):
                    markdown += f"**{stage3['model']}:**\n\n{stage3['response']}\n\n"

            markdown += "---\n\n"

    return {
        "markdown": markdown,
        "filename": f"{conversation['title'].replace(' ', '_')}.md"
    }


@app.get("/api/conversations/{conversation_id}/export/html")
async def export_conversation_html(conversation_id: str):
    """Export a conversation as a self-contained HTML page."""
    conversation = storage.get_conversation(conversation_id)
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found")

    sections = []
    for message in conversation["messages"]:
        if message["role"] == "user":
            sections.append({
                "type": "user",
                "content": message["content"]
            })
        elif message["role"] == "assistant":
            if message.get("mode") == "consensus":
                sections.append({
                    "type": "consensus",
                    "participants": message.get("participants", []),
                    "chairman": message.get("chairman") or {},
                    "rounds": [
                        {
                            "round": r.get("round"),
                            "statements": r.get("statements", []),
                            "votes": r.get("votes") or {},
                            "reached": bool(r.get("reached")),
                            "review": r.get("review"),
                        }
                        for r in message.get("rounds", [])
                    ],
                    "synthesis": message.get("synthesis") or {},
                })
            elif message.get("mode") == "hybrid":
                sections.append({
                    "type": "hybrid",
                    "hybrid_phase1": message.get("hybrid_phase1", []),
                    "hybrid_phase2": message.get("hybrid_phase2", []),
                    "hybrid_phase3": message.get("hybrid_phase3") or {},
                    "hybrid_phase4": message.get("hybrid_phase4") or {},
                })
            else:
                stage1_items = [
                    {"model": r["model"], "response": r["response"]}
                    for r in message.get("stage1", [])
                ]
                stage2_items = [
                    {"model": r["model"], "ranking": r["ranking"]}
                    for r in message.get("stage2", [])
                ]
                stage3 = message.get("stage3") or {}
                sections.append({
                    "type": "assistant",
                    "stage1": stage1_items,
                    "stage2": stage2_items,
                    "stage3": stage3
                })

    import json as _json
    sections_json = _json.dumps(sections)
    title_escaped = conversation["title"].replace('&', '&amp;').replace('"', '&quot;').replace('<', '&lt;').replace('>', '&gt;').replace("'", '&#39;')
    created_at = conversation["created_at"]

    html_before = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>{title_escaped} — LLM Council</title>
  <script src="https://cdnjs.cloudflare.com/ajax/libs/marked/9.1.6/marked.min.js"></script>
  <style>
    *, *::before, *::after {{ box-sizing: border-box; }}
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      background: #f8fafc;
      color: #1e293b;
      margin: 0;
      padding: 24px;
    }}
    .page-header {{
      max-width: 900px;
      margin: 0 auto 32px;
      padding-bottom: 16px;
      border-bottom: 2px solid #e2e8f0;
    }}
    .page-header h1 {{
      font-size: 28px;
      font-weight: 700;
      background: linear-gradient(135deg, #2563eb, #3b82f6);
      -webkit-background-clip: text;
      -webkit-text-fill-color: transparent;
      background-clip: text;
      margin: 0 0 6px;
    }}
    .page-header .meta {{ color: #64748b; font-size: 13px; }}
    .container {{ max-width: 900px; margin: 0 auto; }}
    .message {{ margin-bottom: 28px; border-radius: 12px; overflow: hidden; }}
    .user-message {{
      background: linear-gradient(135deg, #dbeafe, #bfdbfe);
      border: 1px solid #93c5fd;
      padding: 20px 24px;
    }}
    .user-label {{
      font-weight: 700;
      color: #1e40af;
      font-size: 13px;
      text-transform: uppercase;
      letter-spacing: 0.05em;
      margin-bottom: 8px;
    }}
    .user-text {{ color: #1e293b; line-height: 1.7; white-space: pre-wrap; }}
    .file-badge-export {{
      display: inline-flex;
      align-items: center;
      gap: 5px;
      padding: 3px 9px;
      margin-bottom: 8px;
      background: rgba(37,99,235,0.12);
      border: 1px solid rgba(37,99,235,0.25);
      border-radius: 12px;
      font-size: 11px;
      font-weight: 600;
      color: #1e40af;
    }}
    .assistant-message {{
      background: #fff;
      border: 1px solid #e2e8f0;
      box-shadow: 0 2px 12px rgba(0,0,0,0.06);
    }}
    .hybrid-header {{
      padding: 12px 20px;
      background: linear-gradient(135deg, #1e1b4b, #312e81);
      color: white;
      font-weight: 700;
      font-size: 15px;
    }}
    .consensus-header {{
      padding: 12px 20px;
      background: linear-gradient(135deg, #064e3b, #059669);
      color: white;
      font-weight: 700;
      font-size: 15px;
    }}
    .consensus-meta {{
      display: flex;
      gap: 24px;
      flex-wrap: wrap;
      padding: 12px 20px 0;
      font-size: 13px;
      color: #64748b;
    }}
    .consensus-meta strong {{ color: #1e293b; }}
    .vote-tally {{
      display: inline-flex;
      gap: 8px;
      margin-bottom: 14px;
    }}
    .vote-yes, .vote-no {{
      padding: 3px 10px;
      border-radius: 12px;
      font-size: 12px;
      font-weight: 700;
    }}
    .vote-yes {{ background: #d1fae5; color: #065f46; border: 1px solid #6ee7b7; }}
    .vote-no {{ background: #ffe4e6; color: #9f1239; border: 1px solid #fda4af; }}
    .debate-statement {{
      margin-bottom: 16px;
      border: 1px solid #e2e8f0;
      border-radius: 8px;
      overflow: hidden;
    }}
    .debate-statement .model-label {{
      padding: 8px 16px 0;
      margin-bottom: 0;
    }}
    .debate-statement .md-content {{
      border: none;
      border-radius: 0;
    }}
    .position-box {{
      padding: 10px 16px;
      background: #f0fdf4;
      border-top: 1px solid #a7f3d0;
      font-size: 13.5px;
      line-height: 1.6;
      color: #065f46;
    }}
    .review-box {{
      margin-bottom: 14px;
      padding: 10px 14px;
      border-radius: 8px;
      font-size: 13.5px;
      line-height: 1.6;
    }}
    .review-box.aligned {{ background: #eff6ff; border: 1px solid #bfdbfe; color: #1e40af; }}
    .review-box.conflicted {{ background: #fffbeb; border: 1px solid #fde68a; color: #92400e; }}
    .stage-block {{ padding: 20px 24px; border-bottom: 1px solid #f1f5f9; }}
    .stage-block:last-child {{ border-bottom: none; }}
    .stage-heading {{
      font-size: 16px;
      font-weight: 700;
      color: #2563eb;
      margin: 0 0 14px;
      display: flex;
      align-items: center;
      gap: 8px;
    }}
    .stage-heading::before {{
      content: '';
      display: inline-block;
      width: 4px;
      height: 18px;
      background: linear-gradient(180deg, #2563eb, #3b82f6);
      border-radius: 2px;
    }}
    .tabs {{ display: flex; gap: 6px; flex-wrap: wrap; margin-bottom: 14px; }}
    .tab-btn {{
      padding: 7px 14px;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      background: #f8fafc;
      color: #64748b;
      cursor: pointer;
      font-size: 13px;
      font-weight: 500;
      transition: all 0.15s;
    }}
    .tab-btn:hover {{ background: #f1f5f9; color: #2563eb; }}
    .tab-btn.active {{
      background: linear-gradient(135deg, #2563eb, #3b82f6);
      color: #fff;
      border-color: #2563eb;
    }}
    .tab-panel {{ display: none; }}
    .tab-panel.visible {{ display: block; }}
    .md-content {{
      background: #fff;
      border: 1px solid #e2e8f0;
      border-radius: 8px;
      padding: 16px 20px;
      line-height: 1.75;
    }}
    .md-content h1,.md-content h2,.md-content h3 {{
      color: #1e293b; margin-top: 1.2em; margin-bottom: 0.5em;
    }}
    .md-content p {{ margin: 0.6em 0; }}
    .md-content code {{
      background: #f1f5f9; padding: 2px 6px;
      border-radius: 4px; font-size: 0.9em; color: #0f172a;
    }}
    .md-content pre {{
      background: #1e293b; color: #e2e8f0;
      padding: 14px 16px; border-radius: 8px; overflow-x: auto;
    }}
    .md-content pre code {{ background: none; color: inherit; padding: 0; }}
    .md-content blockquote {{
      border-left: 4px solid #3b82f6;
      margin: 0; padding: 8px 16px; color: #475569;
    }}
    .md-content ul, .md-content ol {{ padding-left: 24px; }}
    .md-content table {{ border-collapse: collapse; width: 100%; }}
    .md-content th, .md-content td {{
      border: 1px solid #e2e8f0; padding: 8px 12px; text-align: left;
    }}
    .md-content th {{ background: #f8fafc; font-weight: 600; }}
    .model-label {{
      font-size: 11px; font-family: monospace;
      color: #94a3b8; margin-bottom: 8px;
    }}
    .stage3-block {{
      background: linear-gradient(135deg, #d1fae5, #a7f3d0);
      border-bottom: none;
    }}
    .stage3-block .md-content {{ border-color: #6ee7b7; }}
    .stage3-block .md-content h1,
    .stage3-block .md-content h2,
    .stage3-block .md-content h3 {{ color: #059669; }}
    .aggregate-box {{
      margin-top: 16px;
      background: linear-gradient(135deg, #dbeafe, #bfdbfe);
      border: 1px solid #93c5fd;
      border-radius: 10px;
      padding: 16px 20px;
    }}
    .aggregate-box h4 {{ color: #1e40af; margin: 0 0 10px; font-size: 15px; }}
    .rank-row {{
      display: flex; align-items: center; gap: 12px;
      padding: 10px 14px; margin-bottom: 8px;
      background: #fff; border-radius: 8px;
      border: 1px solid #93c5fd;
    }}
    .rank-pos {{ font-weight: 700; color: #2563eb; font-size: 18px; min-width: 36px; }}
    .rank-model {{ flex: 1; font-size: 14px; font-weight: 500; }}
    .rank-avg {{ font-size: 12px; color: #64748b; }}
  </style>
</head>
<body>
  <div class="page-header">
    <h1>{title_escaped}</h1>
    <div class="meta">Model Behavior by Niiblr &mdash; {created_at}</div>
  </div>
  <div class="container" id="root"></div>

  <script>
    const sections = """

    html_after = """;

    function md(text) {
      return marked.parse(text || '');
    }

    function shortModel(m) {
      return (m || '').split('/')[1] || m;
    }

    let tabCounters = 0;

    function buildTabs(items, labelFn, contentFn) {
      const groupId = 'tg' + (tabCounters++);
      const tabsHtml = items.map((item, i) =>
        `<button class="tab-btn ${i===0?'active':''}" onclick="selectTab('${groupId}',${i},this)">${labelFn(item, i)}</button>`
      ).join('');
      const panelsHtml = items.map((item, i) =>
        `<div class="tab-panel ${i===0?'visible':''}" id="${groupId}-panel-${i}">${contentFn(item, i)}</div>`
      ).join('');
      return `<div class="tabs">${tabsHtml}</div>${panelsHtml}`;
    }

    function selectTab(groupId, idx, btn) {
      const group = btn.closest('.stage-block, .assistant-message');
      group.querySelectorAll(`[id^="${groupId}-panel-"]`).forEach((p,i) => {
        p.classList.toggle('visible', i===idx);
      });
      btn.parentElement.querySelectorAll('.tab-btn').forEach((b,i) => {
        b.classList.toggle('active', i===idx);
      });
    }

    const root = document.getElementById('root');

    sections.forEach((section, si) => {
      const div = document.createElement('div');
      div.className = 'message';

      if (section.type === 'user') {
        div.className += ' user-message';
        // Parse optional file badge from content starting with [File: ...]
        const fileMatch = section.content.match(/^\\[File: (.+?)\\]/);
        const fileName = fileMatch ? fileMatch[1] : null;
        // Strip file block, show only the user question
        let displayText = section.content;
        const questionMatch = section.content.match(/\\nUser question: ([\\s\\S]*)$/);
        if (questionMatch) displayText = questionMatch[1];

        const badgeHtml = fileName
          ? `<div class="file-badge-export">&#128196; ${escHtml(fileName)}</div>`
          : '';
        div.innerHTML = `<div class="user-label">You</div>${badgeHtml}<div class="user-text">${escHtml(displayText)}</div>`;

      } else if (section.type === 'hybrid') {
        div.className += ' assistant-message';

        const phases = [
          { label: '&#x1F4AC; Phase 1: Socratic', items: section.hybrid_phase1, multi: true },
          { label: '&#x2694;&#xFE0F; Phase 2: Debate', items: section.hybrid_phase2, multi: true },
          { label: "&#x1F608; Phase 3: Devils Advocate", single: section.hybrid_phase3, multi: false },
          { label: '&#x2728; Phase 4: Final Synthesis', single: section.hybrid_phase4, multi: false },
        ];

        let html = '<div class="hybrid-header">&#x1F500; Debate Mode &#x2014; Socratic &#x2192; Debate &#x2192; Devils Advocate &#x2192; Synthesis</div>';
        phases.forEach(phase => {
          html += `<div class="stage-block"><div class="stage-heading">${phase.label}</div>`;
          if (phase.multi && phase.items && phase.items.length > 0) {
            html += buildTabs(
              phase.items,
              r => escHtml(r.model),
              r => `<div class="model-label">${escHtml(r.model)}</div><div class="md-content">${md(r.response)}</div>`
            );
          } else if (!phase.multi && phase.single && phase.single.response) {
            html += `<div class="model-label">${escHtml(phase.single.model||'')}</div><div class="md-content">${md(phase.single.response)}</div>`;
          }
          html += '</div>';
        });
        div.innerHTML = html;

      } else if (section.type === 'consensus') {
        div.className += ' assistant-message';

        const participantNames = (section.participants || []).map(p => escHtml(p.name)).join(', ');
        let html = '<div class="consensus-header">&#x1F91D; Consensus Mode &#x2014; Majority-vote debate</div>';
        html += `<div class="consensus-meta">` +
          `<span><strong>Participants:</strong> ${participantNames}</span>` +
          `<span><strong>Chairman:</strong> ${escHtml((section.chairman && section.chairman.name) || '')}</span>` +
          `</div>`;

        (section.rounds || []).forEach(round => {
          const label = round.reached
            ? `Round ${round.round} &#x2014; Majority Consensus`
            : `Round ${round.round}`;
          html += `<div class="stage-block"><div class="stage-heading">${label}</div>`;

          const votes = round.votes || {};
          if ((votes.yes || []).length || (votes.no || []).length) {
            html += `<div class="vote-tally">` +
              `<span class="vote-yes">YES: ${(votes.yes || []).length}</span>` +
              `<span class="vote-no">NO: ${(votes.no || []).length}</span></div>`;
          }

          if (round.review) {
            const cls = round.review.aligned ? 'aligned' : 'conflicted';
            const verdict = round.review.aligned ? 'ALIGNED' : 'CONFLICTED';
            html += `<div class="review-box ${cls}">` +
              `<strong>Chairman review: ${verdict}</strong> &#x2014; ${escHtml(round.review.reasoning || '')}</div>`;
          }

          (round.statements || []).forEach(stmt => {
            html += `<div class="debate-statement">` +
              `<div class="model-label">${escHtml(stmt.model)}</div>` +
              `<div class="md-content">${md(stmt.response)}</div>`;
            if (stmt.position) {
              html += `<div class="position-box"><strong>Final position:</strong> ${md(stmt.position)}</div>`;
            }
            html += `</div>`;
          });

          html += `</div>`;
        });

        const syn = section.synthesis || {};
        html += `<div class="stage-block stage3-block"><div class="stage-heading">Council Consensus</div>` +
          `<div class="model-label">${escHtml(syn.model || '')}</div>` +
          `<div class="md-content">${md(syn.response)}</div></div>`;

        div.innerHTML = html;

      } else {
        div.className += ' assistant-message';

        let s1Html = '';
        if (section.stage1 && section.stage1.length > 0) {
          s1Html = buildTabs(
            section.stage1,
            (r) => shortModel(r.model),
            (r) => `<div class="model-label">${escHtml(r.model)}</div><div class="md-content">${md(r.response)}</div>`
          );
        }

        let s2Html = '';
        if (section.stage2 && section.stage2.length > 0) {
          s2Html = buildTabs(
            section.stage2,
            (r) => shortModel(r.model),
            (r) => `<div class="model-label">${escHtml(r.model)}</div><div class="md-content">${md(r.ranking)}</div>`
          );
        }

        const s3 = section.stage3 || {};
        const s3Html = s3.response
          ? `<div class="model-label">${escHtml(s3.model||'')}</div><div class="md-content">${md(s3.response)}</div>`
          : '<em>Not available</em>';

        div.innerHTML = `
          <div class="stage-block">
            <div class="stage-heading">Stage 1: Individual Responses</div>
            ${s1Html}
          </div>
          <div class="stage-block">
            <div class="stage-heading">Stage 2: Peer Rankings</div>
            ${s2Html}
          </div>
          <div class="stage-block stage3-block">
            <div class="stage-heading">Stage 3: Final Synthesis</div>
            ${s3Html}
          </div>
        `;
      }

      root.appendChild(div);
    });

    function escHtml(str) {
      return String(str)
        .replace(/&/g,'&amp;')
        .replace(/</g,'&lt;')
        .replace(/>/g,'&gt;')
        .replace(/"/g,'&quot;');
    }
  </script>
</body>
</html>"""

    html = html_before + sections_json + html_after

    safe_title = conversation["title"].replace(" ", "_").replace("/", "-")
    return {
        "html": html,
        "filename": f"{safe_title}.html"
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8001)
